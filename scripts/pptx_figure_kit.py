# -*- coding: utf-8 -*-
"""Minimal helpers for academic PPTX figures (python-pptx)."""
from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt, Emu
from lxml import etree

# Default academic palette
PAGE = RGBColor(0xFF, 0xFF, 0xFF)
SYS = RGBColor(0xF2, 0xF2, 0xF0)
SYS_LN = RGBColor(0xB5, 0xB5, 0xAD)
INK = RGBColor(0x28, 0x28, 0x28)
MUTED = RGBColor(0x6E, 0x6E, 0x6E)
GREY = RGBColor(0xB8, 0xB8, 0xB8)
GREY_LN = RGBColor(0x8A, 0x8A, 0x8A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ARROW = RGBColor(0xB2, 0xB2, 0xAA)

M1, M2, M3, M4 = (
    RGBColor(0xBA, 0x9E, 0x76),
    RGBColor(0xBA, 0x70, 0x62),
    RGBColor(0x58, 0x7C, 0x9E),
    RGBColor(0xA8, 0x8E, 0x4E),
)
M1f, M2f, M3f, M4f = (
    RGBColor(0xF8, 0xF3, 0xE9),
    RGBColor(0xF8, 0xEC, 0xE9),
    RGBColor(0xEC, 0xF3, 0xF8),
    RGBColor(0xF7, 0xF3, 0xEA),
)

TITLE_FONT = "Times New Roman"


def new_slide(width_in=13.333, height_in=7.5):
    prs = Presentation()
    prs.slide_width = Inches(width_in)
    prs.slide_height = Inches(height_in)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    _fill(bg, PAGE, None)
    return prs, slide


def _fill(sh, fill, line=None, lw=1.0):
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)


def round_rect(slide, l, t, w, h, fill, line=None, adj=0.06, lw=1.0):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    _fill(sh, fill, line, lw)
    try:
        sh.adjustments[0] = adj
    except Exception:
        pass
    return sh


def ph_image(slide, l, t, w, h):
    """Grey image placeholder (photo slot)."""
    return round_rect(slide, l, t, w, h, GREY, GREY_LN, adj=0.05, lw=0.6)


def label(slide, l, t, w, h, text, size=10, bold=False, color=INK,
          align=PP_ALIGN.CENTER, font=TITLE_FONT, italic=False):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    return box


def task_header(slide, l, t, w, title, question, accent_fill):
    """Title + italic question strip."""
    bar = round_rect(slide, l, t, w, Inches(0.32), accent_fill, None, adj=0.12)
    tf = bar.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.name = TITLE_FONT
    r.font.color.rgb = INK
    label(slide, l, t + Inches(0.34), w, Inches(0.22), question,
          size=9, bold=False, color=MUTED, italic=True)
    return bar


def arrow_right(slide, l, t, w, h, text="", fill=ARROW):
    sh = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, l, t, w, h)
    _fill(sh, fill, RGBColor(0x8A, 0x8A, 0x82), 0.5)
    if text:
        tf = sh.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        r.font.size = Pt(8)
        r.font.bold = True
        r.font.color.rgb = WHITE
        r.font.name = TITLE_FONT
    return sh


def arrow_down(slide, l, t, w, h, text, fill):
    sh = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, l, t, w, h)
    _fill(sh, fill, None, 0.5)
    tf = sh.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(7)
    r.font.bold = True
    r.font.color.rgb = WHITE
    r.font.name = TITLE_FONT
    return sh


def scale_fonts_in_element(el, scale: float):
    """Scale OOXML font sz (100ths of a point) under an element."""
    A = "http://schemas.openxmlformats.org/drawingml/2006/main"
    for node in el.iter():
        if not isinstance(node.tag, str):
            continue
        local = etree.QName(node).localname
        if local in ("rPr", "defRPr", "endParaRPr"):
            sz = node.get("sz")
            if sz and sz.isdigit():
                node.set("sz", str(max(100, int(round(int(sz) * scale)))))
        if local == "ln":
            w = node.get("w")
            if w and w.isdigit():
                node.set("w", str(max(635, int(round(int(w) * scale)))))


def export_png_via_com(pptx_path: str, png_path: str, width_px: int = 2600):
    """Export slide 1 as PNG, without quitting a shared PowerPoint instance."""
    from pathlib import Path
    import win32com.client  # type: ignore

    if not isinstance(width_px, int) or isinstance(width_px, bool) or width_px <= 0:
        raise ValueError("width_px must be a positive integer")
    source = Path(pptx_path).resolve(strict=True)
    target = Path(png_path).resolve()
    if target == source:
        raise ValueError("PNG output must not overwrite the source presentation")
    target.parent.mkdir(parents=True, exist_ok=True)
    app = win32com.client.Dispatch("PowerPoint.Application")
    pres = None
    try:
        pres = app.Presentations.Open(str(source), ReadOnly=True, WithWindow=False)
        sw = float(pres.PageSetup.SlideWidth)
        sh = float(pres.PageSetup.SlideHeight)
        height_px = max(1, int(width_px * sh / sw))
        pres.Slides(1).Export(str(target), "PNG", width_px, height_px)
    finally:
        try:
            if pres is not None:
                pres.Close()
        finally:
            pres = None
            app = None
    return png_path, width_px, height_px
